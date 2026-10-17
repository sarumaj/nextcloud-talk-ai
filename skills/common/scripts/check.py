#!/usr/bin/env python3
"""check.py FILE [--formulas] [--lang de|pl|en] [--open-only]: automatic checks of a file.

Always: the package is intact (zip, well-formed XML, the parts its format needs), the
file opens headless (LibreOffice converts it to PDF), has pages, no page without
content, the fonts it uses, under 20 MB.
Structure (documents): heading levels without jumps, an up-to-date table of contents,
no runs of empty paragraphs, images and tables within the page, alt texts.
Structure (slides): no empty slides, shapes within the slide, text that may overflow,
speaker notes, alt texts. Spreadsheets: formulas (with --formulas they are required),
no error values. With --lang: words Hunspell does not know (a list, not a failure).
ODF files are checked through a temporary OOXML copy.
Images (png jpg webp gif svg): it opens, its size in pixels, not blank; SVG: well-formed,
a viewBox, nothing loaded from outside, the fonts it names are installed.
Exits 1 when a check fails."""
import os, re, shutil, subprocess, sys, tempfile, zipfile

args = sys.argv[1:]
path = args[0] if args else ""
want_formulas = "--formulas" in args
open_only = "--open-only" in args
lang = args[args.index("--lang") + 1] if "--lang" in args and args.index("--lang") + 1 < len(args) else ""
problems, notes = [], []
tmp = tempfile.mkdtemp(dir="/tmp")
lower = path.lower()
ext = os.path.splitext(lower)[1].lstrip(".")


def soffice(target, src, outdir):
    run = subprocess.run(["soffice", "--headless", "--norestore", "--convert-to", target, "--outdir", outdir, src], capture_output=True, text=True, timeout=180)
    out = os.path.join(outdir, os.path.splitext(os.path.basename(src))[0] + "." + target.split(":")[0])
    return out if run.returncode == 0 and os.path.exists(out) else None


def package():
    """The zip container and its XML parts (F-25: a basic validation, not the full schema)."""
    if ext not in ("docx", "xlsx", "pptx", "odt", "ods", "odp", "dotx", "potx", "xltx", "ott", "otp", "ots"):
        return
    try:
        with zipfile.ZipFile(path) as z:
            bad = z.testzip()
            if bad:
                problems.append("damaged zip entry: " + bad)
            names = z.namelist()
            if ext.startswith("o"):
                if names[:1] != ["mimetype"] or z.getinfo("mimetype").compress_type != zipfile.ZIP_STORED:
                    problems.append("ODF package: 'mimetype' must be the first entry, uncompressed")
                for part in ("content.xml", "META-INF/manifest.xml"):
                    if part not in names:
                        problems.append("ODF package lacks " + part)
            elif "[Content_Types].xml" not in names:
                problems.append("OOXML package lacks [Content_Types].xml")
            from lxml import etree
            for name in names:
                if name.endswith((".xml", ".rels")):
                    try:
                        etree.fromstring(z.read(name))
                    except Exception as e:
                        problems.append("malformed XML in %s: %s" % (name, str(e)[:120]))
    except zipfile.BadZipFile:
        problems.append("not a valid zip package")


def opens():
    pdf = path if lower.endswith(".pdf") else soffice("pdf", path, tmp)
    if not pdf:
        problems.append("LibreOffice cannot open it")
        return None
    info = subprocess.run(["pdfinfo", pdf], capture_output=True, text=True).stdout
    pages = int(re.search(r"Pages:\s+(\d+)", info).group(1)) if "Pages:" in info else 0
    notes.append("pages: %d" % pages)
    if pages == 0:
        problems.append("no pages")
    return pdf, pages


def pages_content(pdf, pages):
    for n in range(1, pages + 1):
        text = subprocess.run(["pdftotext", "-f", str(n), "-l", str(n), pdf, "-"], capture_output=True, text=True).stdout
        if not text.strip():
            images = subprocess.run(["pdfimages", "-f", str(n), "-l", str(n), "-list", pdf], capture_output=True, text=True).stdout
            if len(images.splitlines()) <= 2:
                problems.append("page %d is empty" % n)
    fonts = subprocess.run(["pdffonts", pdf], capture_output=True, text=True).stdout.splitlines()[2:]
    used = sorted({re.sub(r"^[A-Z]{6}\+", "", f.split()[0]) for f in fonts if f.split()})
    notes.append("fonts: " + ", ".join(used))
    if lower.endswith(".pdf") and any(len(line.split()) > 5 and line.split()[-5] == "no" for line in fonts):
        problems.append("fonts are not embedded in the PDF")


def ooxml_copy():
    """The file as OOXML for the structure checks (ODF through LibreOffice)."""
    target = {"odt": "docx", "ott": "docx", "odp": "pptx", "otp": "pptx", "ods": "xlsx", "ots": "xlsx"}.get(ext)
    if not target:
        return path
    copy = soffice(target, path, tmp)
    if copy:
        notes.append("structure checked through an OOXML copy")
    return copy


def check_document(docx_path):
    from docx import Document
    from docx.shared import Emu
    from docx.oxml.ns import qn
    doc = Document(docx_path)
    section = doc.sections[0]
    width = (section.page_width or Emu(0)) - (section.left_margin or Emu(0)) - (section.right_margin or Emu(0))
    levels, empty_run, empty_runs, headings = [], 0, 0, []
    for p in doc.paragraphs:
        style = (p.style.name if p.style is not None else "") or ""
        m = re.match(r"(?:Heading|Überschrift|Nagłówek)\s*(\d)", style)
        if m:
            level = int(m.group(1))
            if levels and level > levels[-1] + 1:
                problems.append('heading level jumps from %d to %d at "%s"' % (levels[-1], level, p.text[:50]))
            levels.append(level)
            headings.append(p.text.strip())
        has_break = any(br.get(qn("w:type")) == "page" for br in p._p.iter(qn("w:br")))
        has_drawing = any(True for _ in p._p.iter(qn("w:drawing")))
        if not p.text.strip() and not has_break and not has_drawing:
            empty_run += 1
            if empty_run == 2:
                empty_runs += 1
        else:
            empty_run = 0
    if empty_runs:
        problems.append("%d place(s) with two or more empty paragraphs in a row (use spacing, not empty lines)" % empty_runs)
    body = doc.element.body
    instr = " ".join(t.text or "" for t in body.iter(qn("w:instrText")))
    if "TOC" in instr:
        toc_text = " ".join(t.text or "" for sdt in body.iter(qn("w:sdtContent")) for t in sdt.iter(qn("w:t")))
        if not toc_text.strip() or "No table of contents" in toc_text or "Keine Einträge" in toc_text:
            problems.append("the table of contents is empty or not updated: run python3 /skills/common/scripts/lo_update.py FILE")
        else:
            missing = [h for h in headings[:30] if h and h[:25] not in toc_text]
            if missing:
                problems.append("the table of contents is out of date (missing: %s): run lo_update.py" % "; ".join(m[:30] for m in missing[:3]))
    no_alt = 0
    for shape in doc.inline_shapes:
        if width and shape.width and shape.width > width + Emu(10000):
            problems.append("an image is wider than the text area (%.1f cm > %.1f cm)" % (shape.width.cm, width.cm))
        doc_pr = shape._inline.find(qn("wp:docPr"))
        if doc_pr is not None and not (doc_pr.get("descr") or "").strip():
            no_alt += 1
    if no_alt:
        notes.append("%d image(s) without alt text" % no_alt)
    for i, table in enumerate(doc.tables, 1):
        grid = table._tbl.find(qn("w:tblGrid"))
        total = sum(int(c.get(qn("w:w")) or 0) for c in grid.iter(qn("w:gridCol"))) * 635 if grid is not None else 0
        # Cell widths of the widest row win over the grid (python-docx sets those).
        for row in table._tbl.iter(qn("w:tr")):
            cells = sum(int(w.get(qn("w:w")) or 0) for w in row.iter(qn("w:tcW")) if w.get(qn("w:type"), "dxa") == "dxa") * 635
            total = max(total, cells)
        if width and total > width * 1.02:
            problems.append("table %d is wider than the text area" % i)
    notes.append("headings: %d, tables: %d, images: %d" % (len(headings), len(doc.tables), len(doc.inline_shapes)))


def check_slides(pptx_path):
    from pptx import Presentation
    from pptx.util import Pt
    prs = Presentation(pptx_path)
    w, h = prs.slide_width, prs.slide_height
    no_notes, no_alt = 0, 0
    for n, slide in enumerate(prs.slides, 1):
        texts = [s.text_frame.text for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
        pictures = [s for s in slide.shapes if s.shape_type == 13 or getattr(s, "has_chart", False) or getattr(s, "has_table", False)]
        if not texts and not pictures:
            problems.append("slide %d is empty" % n)
        for s in slide.shapes:
            if s.left is not None and s.width is not None and (s.left < -Pt(2) or s.top < -Pt(2) or s.left + s.width > w + Pt(2) or s.top + s.height > h + Pt(2)):
                problems.append('slide %d: "%s" reaches beyond the slide' % (n, s.name))
            if s.has_text_frame and s.width and s.height:
                chars = len(s.text_frame.text)
                sizes = [r.font.size.pt for p in s.text_frame.paragraphs for r in p.runs if r.font.size]
                size = min(sizes) if sizes else 24
                capacity = (s.width / 12700) * (s.height / 12700) / (size * size * 0.62)
                if chars > capacity * 1.1 and chars > 40:
                    problems.append('slide %d: the text in "%s" may overflow (%d characters at %.0f pt)' % (n, s.name, chars, size))
            if s.shape_type == 13:
                descr = s._element.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr")
                if descr is not None and not (descr.get("descr") or "").strip():
                    no_alt += 1
        if not (slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip()):
            no_notes += 1
    notes.append("slides: %d" % len(prs.slides))
    if no_notes:
        notes.append("%d slide(s) without speaker notes" % no_notes)
    if no_alt:
        notes.append("%d picture(s) without alt text" % no_alt)


def check_sheet(xlsx_path):
    with zipfile.ZipFile(xlsx_path) as z:
        sheets = [z.read(n).decode("utf-8", "replace") for n in z.namelist() if n.startswith("xl/worksheets/sheet")]
    formulas = sum(s.count("<f>") + s.count("<f ") for s in sheets)
    errors = sum(len(re.findall(r"<v>#(REF|DIV/0|VALUE|NAME|N/A|NUM)!?", s)) for s in sheets)
    notes.append("formulas: %d" % formulas)
    if want_formulas and formulas == 0:
        problems.append("no formulas, but calculations were asked for (write =SUM(…) etc., not results)")
    if errors:
        problems.append("%d cells show errors (#REF!, #DIV/0! …)" % errors)


def spelling(pdf):
    dictionary = {"de": "de_DE", "pl": "pl_PL", "en": "en_US"}.get(lang)
    if not dictionary or not pdf or not shutil.which("hunspell"):
        return
    text = subprocess.run(["pdftotext", pdf, "-"], capture_output=True, text=True).stdout
    run = subprocess.run(["hunspell", "-l", "-d", dictionary], input=text, capture_output=True, text=True)
    words = [w for w in dict.fromkeys(run.stdout.split()) if len(w) > 2 and not w.isupper() and not re.search(r"\d", w)]
    if words:
        notes.append("unknown to the %s dictionary (check, names may be fine): %s" % (dictionary, ", ".join(words[:30])))


def check_image():
    """An image opens and is not blank; an SVG stands alone (no outside files or fonts)."""
    png = os.path.join(tmp, "image.png")
    if ext == "svg":
        from lxml import etree
        try:
            root = etree.parse(path).getroot()
        except Exception as e:
            problems.append("malformed SVG: %s" % str(e)[:160])
            return
        if root.get("viewBox") is None:
            notes.append("no viewBox: the SVG does not scale cleanly")
        for el in root.iter():
            href = el.get("href") or el.get("{http://www.w3.org/1999/xlink}href") or ""
            if href and not href.startswith(("#", "data:")):
                problems.append("loads %s from outside: embed it (data: URI) or draw it" % href[:80])
        text = etree.tostring(root, encoding="unicode")
        if re.search(r"@import|url\((?!#|['\"]?#|['\"]?data:)", text):
            problems.append("CSS loads something from outside (@import or url()): embed it")
        installed = subprocess.run(["fc-list", ":", "family"], capture_output=True, text=True).stdout.lower()
        named = set(re.findall(r"font-family\s*[:=]\s*[\"']?([^;\"'>]+)", text))
        generic = {"serif", "sans-serif", "monospace", "cursive", "fantasy", "system-ui"}
        for family in sorted(named):
            first = family.split(",")[0].strip().strip("'\"")
            if first and first.lower() not in generic and first.lower() not in installed:
                problems.append("font %r is not installed and would be replaced (see the fonts in common/SKILL.md)" % first)
        run = subprocess.run(["rsvg-convert", "-w", "800", "--keep-aspect-ratio", path, "-o", png], capture_output=True, text=True)
        if run.returncode != 0:
            problems.append("the SVG does not render: " + run.stderr.strip()[:200])
            return
    else:
        run = subprocess.run(["convert", path + "[0]", png], capture_output=True, text=True)
        if run.returncode != 0:
            problems.append("the image does not open: " + run.stderr.strip()[:200])
            return
    info = subprocess.run(["identify", "-format", "%w %h",
                           path + "[0]" if ext != "svg" else png], capture_output=True, text=True).stdout.split()
    if ext == "svg":
        box = (root.get("viewBox") or "").replace(",", " ").split()
        width, height = (re.fullmatch(r"([\d.]+)(px)?", root.get(k) or "") for k in ("width", "height"))
        info = [width.group(1), height.group(1)] if width and height else box[2:4] if len(box) == 4 else []
    if len(info) >= 2:
        w, h = round(float(info[0])), round(float(info[1]))
        notes.append("size: %d × %d px%s" % (w, h, " (a vector image: scales freely)" if ext == "svg" else ""))
        if ext != "svg" and max(w, h) < 400:
            problems.append("only %d × %d px: too small to view; make it larger (e.g. 1600 px wide)" % (w, h))
        if ext != "svg" and max(w, h) > 8000:
            problems.append("%d × %d px is needlessly large: at most about 4000 px" % (w, h))
    flat = subprocess.run(["convert", png, "-background", "white", "-alpha", "remove", "-alpha", "off", "-format", "%[fx:standard_deviation]", "info:"],
                          capture_output=True, text=True).stdout.strip()
    try:
        if float(flat) < 0.003:
            problems.append("the image is blank (one colour)")
    except ValueError:
        pass


print("checks for", path)
if not os.path.isfile(path):
    print("  ✗ no such file")
    sys.exit(1)
size = os.path.getsize(path)
if size > 20 * 1024 * 1024:
    problems.append("larger than 20 MB (%.1f MB)" % (size / 1048576))
image = ext in ("png", "jpg", "jpeg", "webp", "gif", "svg")
if image:
    check_image()
package()
opened = None if image else opens()
if not open_only and opened:
    pages_content(*opened)
    try:
        copy = ooxml_copy()
        if copy and copy.endswith(("docx", "dotx")):
            check_document(copy)
        elif copy and copy.endswith(("pptx", "potx")):
            check_slides(copy)
        elif copy and copy.endswith(("xlsx", "xltx")):
            check_sheet(copy)
            if ext == "ods":
                with zipfile.ZipFile(path) as z:
                    content = z.read("content.xml").decode("utf-8", "replace")
                if want_formulas and content.count("table:formula=") == 0:
                    problems.append("no formulas in the ODS")
    except Exception as e:
        notes.append("structure check skipped: %s" % str(e)[:200])
    spelling(opened[0])
for n in notes:
    print("  ·", n)
for p in problems:
    print("  ✗", p)
print("  ✓ all checks passed" if not problems else "  %d problem(s)" % len(problems))
shutil.rmtree(tmp, ignore_errors=True)
sys.exit(1 if problems else 0)
