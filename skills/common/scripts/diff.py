#!/usr/bin/env python3
"""diff.py ORIGINAL CHANGED: what changed between two versions of a document.

Compares the text (paragraph by paragraph, with styles), and reports whether styles,
headers and footers, comments, tracked changes and embedded objects were kept. For
slides: titles and text per slide; for spreadsheets: values and formulas per sheet.
ODF files are compared through temporary OOXML copies. The model sees this after each
render of a changed document; it goes into the job's protocol."""
import difflib, os, re, subprocess, sys, tempfile, zipfile

original, changed = sys.argv[1], sys.argv[2]
tmp = tempfile.mkdtemp(dir="/tmp")


def ooxml(path):
    ext = os.path.splitext(path)[1].lower().lstrip(".")
    target = {"odt": "docx", "ott": "docx", "odp": "pptx", "otp": "pptx", "ods": "xlsx", "ots": "xlsx", "doc": "docx", "ppt": "pptx", "xls": "xlsx"}.get(ext)
    if not target:
        return path
    out = os.path.join(tmp, os.path.basename(os.path.dirname(path)) + "-" + os.path.basename(path))
    os.makedirs(out, exist_ok=True)
    subprocess.run(["soffice", "--headless", "--norestore", "--convert-to", target, "--outdir", out, path], capture_output=True, timeout=180)
    return os.path.join(out, os.path.splitext(os.path.basename(path))[0] + "." + target)


def parts(path):
    """Counts of what must survive a change, from the package."""
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        read = lambda n: z.read(n).decode("utf-8", "replace")
        xml = "".join(read(n) for n in names if n.endswith(".xml") and ("document" in n or "slide" in n or "sheet" in n))
        styles = read("word/styles.xml") if "word/styles.xml" in names else ""
        hf = " ".join(re.sub(r"<[^>]+>", " ", read(n)) for n in names if re.match(r"word/(header|footer)\d*\.xml", n))
        comments = "".join(read(n) for n in names if re.search(r"comments\d*\.xml$", n))
    return {
        "styles": set(re.findall(r'w:styleId="([^"]+)"', styles)),
        "headers and footers": " ".join(hf.split()),
        "comments": len(re.findall(r"<w:comment |<p:cm ", comments)),
        "tracked changes": len(re.findall(r"<w:(ins|del) ", xml)),
        "embedded objects": len(re.findall(r"<w:object|<p:oleObj|<c:chart ", xml)) + sum(1 for n in names if n.startswith(("word/embeddings/", "ppt/embeddings/", "xl/embeddings/"))),
        "images": sum(1 for n in names if "/media/" in n),
    }


def lines(path):
    ext = os.path.splitext(path)[1].lower()
    if ext in (".docx", ".dotx"):
        from docx import Document
        doc = Document(path)
        out = ["[%s] %s" % (p.style.name if p.style is not None else "", p.text) for p in doc.paragraphs if p.text.strip()]
        for i, t in enumerate(doc.tables, 1):
            out += ["[table %d] %s" % (i, " | ".join(c.text for c in row.cells)) for row in t.rows]
        return out
    if ext in (".pptx", ".potx"):
        from pptx import Presentation
        out = []
        for n, slide in enumerate(Presentation(path).slides, 1):
            for s in slide.shapes:
                if s.has_text_frame and s.text_frame.text.strip():
                    out.append("[slide %d] %s" % (n, " / ".join(s.text_frame.text.split("\n"))))
        return out
    if ext in (".xlsx", ".xltx"):
        from openpyxl import load_workbook
        out = []
        for ws in load_workbook(path).worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if c.value is not None:
                        out.append("[%s!%s] %s" % (ws.title, c.coordinate, c.value))
        return out
    return []


a, b = ooxml(original), ooxml(changed)
if not (os.path.exists(a) and os.path.exists(b)):
    print("cannot compare (a version does not open)")
    sys.exit(1)
la, lb = lines(a), lines(b)
diff = list(difflib.unified_diff(la, lb, "original", "changed", n=0, lineterm=""))
added = sum(1 for d in diff if d.startswith("+") and not d.startswith("+++"))
removed = sum(1 for d in diff if d.startswith("-") and not d.startswith("---"))
print("text: %d line(s) added, %d removed or changed, %d unchanged" % (added, removed, len(la) - removed))
try:
    pa, pb = parts(a), parts(b)
    for key in pa:
        if key == "styles":
            lost = pa[key] - pb[key]
            print("styles: kept" if not lost else "styles: LOST %s" % ", ".join(sorted(lost)[:10]))
        elif pa[key] == pb[key]:
            print("%s: kept" % key)
        else:
            print("%s: CHANGED (%s → %s)" % (key, str(pa[key])[:80], str(pb[key])[:80]))
except Exception as e:
    print("package comparison skipped: %s" % str(e)[:100])
print("--- the changes (- original, + changed):")
for d in diff[2:200]:
    if not d.startswith("@@"):
        print(d[:300])
if len(diff) > 200:
    print("[… %d more lines]" % (len(diff) - 200))
