#!/usr/bin/env python3
"""describe.py FILE (.xlsx .xlsm .ods .csv): sheets, size, the first rows, formulas."""
import csv, subprocess, sys, tempfile, os
path = sys.argv[1]
if path.lower().endswith(".csv"):
    sample = open(path, encoding="utf-8-sig", newline="").read(4096)
    dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
    rows = list(csv.reader(open(path, encoding="utf-8-sig", newline=""), dialect))
    print(f"CSV, delimiter {dialect.delimiter!r}, {len(rows)} rows × {max(map(len, rows)) if rows else 0} columns")
    for r in rows[:8]:
        print("  ", r)
    sys.exit(0)
if path.lower().endswith(".ods"):
    out = tempfile.mkdtemp(dir="/tmp")
    subprocess.run(["soffice", "--headless", "--norestore", "--convert-to", "xlsx", "--outdir", out, path], check=True, capture_output=True)
    path = os.path.join(out, os.path.splitext(os.path.basename(path))[0] + ".xlsx")
from openpyxl import load_workbook
wb = load_workbook(path)
for ws in wb.worksheets:
    formulas = sum(1 for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("="))
    print(f"sheet {ws.title!r}: {ws.max_row} rows × {ws.max_column} columns, {formulas} formulas, {len(ws._charts)} charts, merged {list(ws.merged_cells.ranges)[:5]}")
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 6), values_only=True):
        print("   ", row[:12])
