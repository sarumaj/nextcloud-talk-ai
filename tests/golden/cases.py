#!/usr/bin/env python3
"""The golden cases: each builds a file as a model should (or as it should not) and
states what the checks must say. T-01 ODT protocol, T-02 template, T-03 CSV to XLSX
with formulas and chart, T-04 PPTX deck with notes, T-05 the same as ODP, T-06 a
change that keeps the rest, T-07 PDF export with embedded fonts, T-08 an SVG poster
with fitted text, T-09 a Graphviz flowchart as PNG; and broken files the checks must
catch (heading jump, empty paragraphs, overflowing table, empty slide, typed results
instead of formulas, an image that loads from outside or names a missing font, a blank
or tiny image)."""
import csv, os, subprocess, sys

os.makedirs("/work/out", exist_ok=True)
os.environ["HOME"] = "/tmp"
results = []


def check(path, *flags):
    run = subprocess.run(["python3", "/skills/common/scripts/check.py", path, *flags], capture_output=True, text=True, timeout=300)
    return run.returncode == 0, run.stdout


def soffice(target, path):
    subprocess.run(["soffice", "--headless", "--norestore", "--convert-to", target, "--outdir", "/work/out", path], capture_output=True, timeout=300)
    return os.path.splitext(path)[0] + "." + target.split(":")[0]


def expect(name, ok, report, want=True, must=()):
    good = ok == want and all(m in report for m in must)
    results.append(good)
    print(("PASS " if good else "FAIL ") + name)
    if not good:
        print(report)


from docx import Document
from docx.shared import Cm


def protocol(path, base=None):
    d = Document(base) if base else Document()
    d.add_heading("Protokoll der Vorstandssitzung", 0)
    d.add_paragraph("14.10.2026, 19:00–20:30, Vereinsheim")
    d.add_heading("Tagesordnung", 1)
    for item in ["Kassenbericht", "Sommerfest 2027", "Gartengeräte", "Verschiedenes"]:
        d.add_paragraph(item, style="List Number")
    d.add_heading("Beschlüsse", 1)
    d.add_paragraph("Budget für das Sommerfest: 1.500 EUR (4 dafür, 0 dagegen).")
    d.save(path)


# T-01: a protocol as ODT
protocol("/work/out/t01.docx")
odt = soffice("odt", "/work/out/t01.docx")
ok, report = check(odt, "--lang", "de")
expect("T-01 protocol as ODT opens, has headings, no empty page", ok, report)

# T-02: on a template: header and footer survive
base = Document()
base.sections[0].header.paragraphs[0].text = "Verein Sonnenhof e.V."
base.sections[0].footer.paragraphs[0].text = "Vertraulich"
base.save("/work/out/template.docx")
protocol("/work/out/t02.docx", "/work/out/template.docx")
ok, report = check("/work/out/t02.docx")
diff = subprocess.run(["python3", "/skills/common/scripts/diff.py", "/work/out/template.docx", "/work/out/t02.docx"], capture_output=True, text=True).stdout
expect("T-02 template keeps header and footer", ok, report + diff, must=())
expect("T-02 diff sees header and footer kept", "headers and footers: kept" in diff, diff)

# T-03: CSV to XLSX with formulas and chart
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
rows = list(csv.reader(open("/work/umsatz.csv", encoding="utf-8"), delimiter=";"))
wb = Workbook(); ws = wb.active; ws.append(rows[0])
for r in rows[1:]:
    ws.append([r[0], float(r[1].replace(".", "").replace(",", ".")), float(r[2].replace(".", "").replace(",", "."))])
ws.append(["Summe", "=SUM(B2:B4)", "=SUM(C2:C4)"])
chart = BarChart(); chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=4, max_col=3), titles_from_data=True)
chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=4)); ws.add_chart(chart, "E2")
wb.save("/work/out/t03.xlsx")
ok, report = check("/work/out/t03.xlsx", "--formulas")
expect("T-03 XLSX with formulas", ok, report, must=("formulas: 2",))
wb2 = Workbook(); ws2 = wb2.active; ws2.append(["Summe", 39800.5]); wb2.save("/work/out/typed.xlsx")
ok, report = check("/work/out/typed.xlsx", "--formulas")
expect("caught: typed results instead of formulas", ok, report, want=False, must=("no formulas",))

# T-04 and T-05: a deck with notes, and as ODP
from pptx import Presentation
p = Presentation()
for i in range(10):
    s = p.slides.add_slide(p.slide_layouts[1])
    s.shapes.title.text = "Wasserkreislauf %d" % (i + 1)
    s.placeholders[1].text = "Verdunstung\nKondensation\nNiederschlag"
    s.notes_slide.notes_text_frame.text = "Hinweise für Folie %d" % (i + 1)
p.save("/work/out/t04.pptx")
ok, report = check("/work/out/t04.pptx")
expect("T-04 ten slides, notes, no overflow", ok and "without speaker notes" not in report, report, must=("slides: 10",))
odp = soffice("odp", "/work/out/t04.pptx")
ok, report = check(odp)
expect("T-05 the deck as ODP opens with ten slides", ok, report, must=("pages: 10",))

# T-06: a chapter added, the rest kept
d = Document("/work/out/t02.docx"); d.add_heading("Nächste Schritte", 1); d.add_paragraph("Angebot einholen bis 31.10."); d.save("/work/out/t06.docx")
diff = subprocess.run(["python3", "/skills/common/scripts/diff.py", "/work/out/t02.docx", "/work/out/t06.docx"], capture_output=True, text=True).stdout
expect("T-06 only the new chapter changed", "2 line(s) added, 0 removed" in diff and "styles: kept" in diff, diff)

# T-07: PDF export with embedded fonts and the same pages
pdf = soffice("pdf", "/work/out/t02.docx")
ok, report = check(pdf)
expect("T-07 PDF with embedded fonts", ok and "not embedded" not in report, report)

# Broken files the checks must catch
d = Document(); d.add_heading("A", 1); d.add_heading("C", 3); d.save("/work/out/jump.docx")
ok, report = check("/work/out/jump.docx")
expect("caught: heading jump", ok, report, want=False, must=("jumps from 1 to 3",))
d = Document(); d.add_paragraph("x"); [d.add_paragraph("") for _ in range(3)]; d.add_paragraph("y"); d.save("/work/out/empty.docx")
ok, report = check("/work/out/empty.docx")
expect("caught: empty paragraphs", ok, report, want=False, must=("empty paragraphs",))
d = Document(); t = d.add_table(rows=1, cols=8)
for c in t.rows[0].cells:
    c.width = Cm(4)
d.save("/work/out/wide.docx")
ok, report = check("/work/out/wide.docx")
expect("caught: table wider than the page", ok, report, want=False, must=("wider than the text area",))
p = Presentation(); p.slides.add_slide(p.slide_layouts[6]); p.save("/work/out/empty.pptx")
ok, report = check("/work/out/empty.pptx")
expect("caught: empty slide", ok, report, want=False, must=("slide 1 is empty",))

# T-08: an SVG poster, text fitted with the real font, rendered as the jobs do
sys.path.insert(0, "/skills/image/scripts")
from textfit import fit, svg_text, contrast
size, lines = fit("Sommerfest im Vereinsheim Sonnenhof", "Open Sans", 1300, 260, largest=110, bold=True)
title, _ = svg_text("\n".join(lines), 800, 300, "Open Sans", size, 1300, anchor="middle", fill="#ffffff", bold=True)
when, _ = svg_text("Samstag, 12. Juni 2027, ab 15 Uhr", 800, 700, "Open Sans", 48, 1300, anchor="middle", fill="#ffffff")
svg = ('<svg xmlns="http://www.w3.org/2000/svg" width="1600" height="900" viewBox="0 0 1600 900"><title>Plakat Sommerfest</title>'
       '<rect width="100%" height="100%" fill="#0072B2"/><circle cx="1400" cy="150" r="90" fill="#E69F00"/>' + title + when + '</svg>')
open("/work/out/t08.svg", "w").write(svg)
ok, report = check("/work/out/t08.svg")
render = subprocess.run(["bash", "/skills/common/scripts/render.sh", "/work/out/t08.svg"], capture_output=True, text=True).stdout
expect("T-08 SVG poster: checks pass, renders one page, text has contrast", ok and contrast("#ffffff", "#0072B2") >= 4.5
       and "pages: 1" in render and "sheet: " in render, report + render, must=("size: 1600 × 900",))

# T-09: a flowchart with Graphviz, delivered as PNG
open("/work/out/t09.dot", "w").write('digraph { graph [rankdir=TB, fontname="Open Sans"]; node [shape=box, style="rounded,filled", '
                                     'fillcolor="#e8f1fa", fontname="Open Sans"]; Start [shape=oval]; Start -> Lesen -> '
                                     'Entscheiden; Entscheiden [shape=diamond]; Entscheiden -> Ende [label="ja"]; Entscheiden -> Lesen [label="nein"]; }')
subprocess.run(["dot", "-Tsvg", "/work/out/t09.dot", "-o", "/work/out/t09.svg"], check=True)
subprocess.run(["rsvg-convert", "-w", "1200", "--keep-aspect-ratio", "-b", "white", "/work/out/t09.svg", "-o", "/work/out/t09.png"], check=True)
ok, report = check("/work/out/t09.png")
expect("T-09 Graphviz flowchart as PNG", ok, report, must=("size: 1200 ×",))

# Broken images the checks must catch
open("/work/out/remote.svg", "w").write('<svg xmlns="http://www.w3.org/2000/svg" xmlns:xlink="http://www.w3.org/1999/xlink" width="800" '
                                        'height="600" viewBox="0 0 800 600"><rect width="800" height="600" fill="#fff"/><image '
                                        'xlink:href="https://example.org/cat.png" width="400" height="300"/><text x="10" y="50" '
                                        'font-family="Comic Sans MS" font-size="30">Hallo</text></svg>')
ok, report = check("/work/out/remote.svg")
expect("caught: an SVG that loads from outside and names a missing font", ok, report, want=False,
       must=("from outside", "'Comic Sans MS' is not installed"))
subprocess.run(["convert", "-size", "1600x900", "xc:white", "/work/out/blank.png"], check=True)
ok, report = check("/work/out/blank.png")
expect("caught: a blank image", ok, report, want=False, must=("blank",))
subprocess.run(["convert", "-size", "120x80", "gradient:", "/work/out/tiny.png"], check=True)
ok, report = check("/work/out/tiny.png")
expect("caught: a tiny image", ok, report, want=False, must=("too small",))

print("%d of %d golden cases pass" % (sum(results), len(results)))
sys.exit(0 if all(results) else 1)
